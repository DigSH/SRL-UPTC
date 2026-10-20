"""Genera assets/data.js a partir de Nivel SRL.html y del libro Excel de HELIX.

Uso:
    python tools/build_data.py ["ruta/Nivel SRL.html"] ["ruta/Copia_de_Helix_221004_pro (1).xlsm"]

Mantiene la plataforma SRL-UPTC (SocialTech) sincronizada con las fuentes originales:
1. Cuestionario SRL (Niveles 1-9, 45 preguntas) desde ../Nivel SRL.html.
2. Matriz de Innovación Social HELIX (5 vectores, 50 elementos, 250 frases de rúbrica,
   50 explicaciones metodológicas, 19 ámbitos predeterminados y 8 criterios de preferencias)
   desde ../../HELIX/Copia_de_Helix_221004_pro (1).xlsm.
"""

from html.parser import HTMLParser
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SRL_HTML = ROOT.parent / "Nivel SRL.html"
DEFAULT_HELIX_XLSM = ROOT.parent.parent / "HELIX" / "Copia_de_Helix_221004_pro (1).xlsm"
OUT = ROOT / "assets" / "data.js"

# Erratas evidentes de los archivos originales que se corrigen al generar los datos.
ERRATA = {
    "signitivativa": "significativa",
    "un un producto": "un producto",
    "herramieneta": "herramienta",
    "poderación": "ponderación",
    "Ponderación entrer elementos": "Ponderación entre elementos",
}

# Etapas de preparación social asociadas a cada nivel SRL (1 a 9).
SRL_STAGES = {
    1: "Identificación y diagnóstico",
    2: "Formulación y diseño",
    3: "Formulación y diseño",
    4: "Piloto y validación social",
    5: "Piloto y validación social",
    6: "Implementación y optimización",
    7: "Implementación y optimización",
    8: "Integración y sostenibilidad",
    9: "Integración y sostenibilidad",
}

SRL_MATURITY = {
    0: "Sin nivel consolidado (en diagnóstico inicial)",
    1: "Problema social identificado",
    2: "Soluciones formuladas",
    3: "Prueba conceptual social",
    4: "Prueba piloto en entorno relevante",
    5: "Solución validada con actores",
    6: "Implementación inicial en entorno real",
    7: "Solución optimizada y monitoreada",
    8: "Integración comunitaria e institucional",
    9: "Sostenibilidad y replicabilidad demostrada",
}

# Acciones concretas de la ruta de maduración por pregunta SRL (45 ítems: 9 niveles x 5 preguntas).
SRL_ACTIONS = [
    # Nivel 1
    "Documentar el problema social que aborda el proyecto con evidencias verificables y fuentes primarias o secundarias.",
    "Caracterizar cómo afecta el problema a las comunidades objetivo y dimensionar su magnitud en el territorio.",
    "Levantar indicadores iniciales sobre la disposición e interés de la comunidad para abordar el problema.",
    "Identificar y registrar los grupos sociales clave (líderes, organizaciones sociales, instituciones) potenciales aliados.",
    "Evaluar el nivel de conocimiento, percepción y priorización que tiene la comunidad frente al problema identificado.",
    # Nivel 2
    "Formular de manera precisa, consensuada y comprensible el problema social que se busca transformar.",
    "Identificar y comparar alternativas de solución potenciales frente al problema social formulado.",
    "Reunir evidencias preliminares que respalden el impacto social positivo esperado de la solución.",
    "Elaborar el mapeo de actores clave especificando sus roles, intereses y grado de participación.",
    "Analizar las barreras sociales, culturales, normativas o económicas que pueden condicionar la solución.",
    # Nivel 3
    "Desarrollar una versión inicial o prueba conceptual de la solución social para su discusión temprana.",
    "Involucrar activamente a las partes interesadas relevantes en las sesiones de validación inicial.",
    "Recoger y sistematizar evidencias de aceptación preliminar por parte de la comunidad objetivo.",
    "Documentar las lecciones aprendidas de la validación conceptual y los ajustes requeridos en el diseño.",
    "Definir un plan de trabajo para pasar a prueba piloto incorporando la retroalimentación de la comunidad.",
    # Nivel 4
    "Ejecutar una prueba piloto de la solución en un entorno real o socialmente relevante con usuarios reales.",
    "Recolectar datos estructurados sobre el impacto social potencial durante el desarrollo de la prueba piloto.",
    "Vincular a las partes interesadas relevantes en la evaluación participativa de la prueba piloto.",
    "Identificar y registrar las barreras operativas, culturales o de adopción observadas en el piloto.",
    "Formular planes específicos de ajuste de la solución a partir de los hallazgos de la prueba piloto.",
    # Nivel 5
    "Validar integralmente la solución con las partes interesadas clave tras los resultados de la prueba piloto.",
    "Definir y medir indicadores que demuestren el cumplimiento del impacto social esperado.",
    "Documentar los beneficios tangibles y resultados positivos verificados en la población participante.",
    "Identificar y priorizar las mejoras necesarias para optimizar la operación y apropiación de la solución.",
    "Confirmar con la comunidad y los aliados que la solución es aceptada como opción viable y sostenible.",
    # Nivel 6
    "Implementar la solución en un entorno real con participación activa de los actores comunitarios e institucionales.",
    "Consolidar registros y datos medibles que demuestren el impacto social inicial de la implementación.",
    "Realizar un análisis cualitativo y cuantitativo de la retroalimentación recibida de los participantes.",
    "Identificar y documentar casos de éxito tempranos derivados de la aplicación en entorno real.",
    "Aplicar ajustes operativos sobre la solución a partir de la retroalimentación del entorno real.",
    # Nivel 7
    "Completar los ajustes de optimización técnica, social y metodológica de la solución en operación.",
    "Validar nuevamente la solución optimizada con la comunidad objetivo para verificar su pertinencia.",
    "Establecer procedimientos permanentes para monitorear y evaluar el impacto social a largo plazo.",
    "Documentar un modelo operativo actualizado, viable y transferible a los actores locales.",
    "Comprobar el desempeño de la solución optimizada en condiciones reales del entorno relevante.",
    # Nivel 8
    "Desarrollar e implementar un plan integral para integrar la solución en la dinámica cotidiana de la comunidad.",
    "Capacitar y transferir capacidades a la comunidad para mantener la solución con mínima asistencia externa.",
    "Formalizar alianzas estratégicas con actores locales, entidades públicas o comunitarias que den soporte.",
    "Gestionar lineamientos, políticas o acuerdos normativos e institucionales que respalden la adopción.",
    "Incluir en el plan de integración indicadores de seguimiento y cronograma periódico de evaluación.",
    # Nivel 9
    "Demostrar con evidencia longitudinal la efectividad y sostenibilidad de la solución en el entorno real.",
    "Sistematizar indicadores y protocolos que respalden la replicabilidad del modelo en otras comunidades.",
    "Verificar que la solución opera de manera autónoma y apropiada por parte de la comunidad beneficiaria.",
    "Publicar o documentar historias de éxito, sistematización de experiencias y lecciones aprendidas.",
    "Dejar instalados mecanismos de monitoreo y actualización continua de la solución hacia el futuro.",
]

SRL_LEVEL_DESCRIPTIONS = {
    1: "El proyecto se encuentra en la fase inicial de identificación, delimitación y comprensión documentada del problema social junto con los actores del territorio.",
    2: "Se formulan alternativas de solución social, se mapean los actores clave con sus roles y se analizan las barreras sociales, culturales y económicas.",
    3: "Se cuenta con una versión inicial o prueba conceptual de la solución y se adelanta su validación temprana con las partes interesadas y la comunidad objetivo.",
    4: "Se ejecuta una prueba piloto en un entorno relevante, recolectando evidencias de impacto potencial y detectando barreras de adopción en terreno.",
    5: "La solución es validada con las partes interesadas mediante métricas de impacto social, documentando beneficios tangibles y condiciones de viabilidad.",
    6: "La solución se implementa en un entorno real con participación activa de los actores locales, midiendo el impacto inicial y ajustando su operación.",
    7: "Se optimiza la solución a partir de la evidencia en terreno, se consolida un modelo operativo viable y se instalan procedimientos de monitoreo a largo plazo.",
    8: "La solución se integra en la vida cotidiana e institucional de la comunidad, respaldada por capacidades locales, alianzas estratégicas y soporte normativo.",
    9: "La solución demuestra efectividad y sostenibilidad autónoma en el entorno real y cuenta con condiciones verificadas para su replicabilidad en otros contextos.",
}

SRL_RECOMMENDATIONS = {
    1: [
        "Documente exhaustivamente el problema social con datos y evidencias verificables.",
        "Realice un diagnóstico participativo con la comunidad afectada.",
        "Identifique y mapee a todos los actores clave relevantes en el territorio.",
    ],
    2: [
        "Evalúe múltiples alternativas de solución antes de seleccionar la definitiva.",
        "Analice las barreras sociales, culturales y económicas potenciales.",
        "Establezca métricas preliminares para medir el impacto social esperado.",
    ],
    3: [
        "Desarrolle un prototipo social o versión inicial de la solución.",
        "Involucre a las partes interesadas en la validación temprana.",
        "Documente las lecciones aprendidas durante el desarrollo conceptual.",
    ],
    4: [
        "Diseñe y ejecute una prueba piloto con metodología rigurosa y participativa.",
        "Establezca mecanismos para recoger datos cualitativos y cuantitativos durante la prueba.",
        "Prepare planes de contingencia frente a las barreras identificadas en terreno.",
    ],
    5: [
        "Valide la solución con métricas cuantitativas y cualitativas de impacto social.",
        "Documente beneficios tangibles y resultados positivos en la comunidad.",
        "Identifique áreas de mejora para optimizar la solución antes del escalamiento.",
    ],
    6: [
        "Implemente la solución con participación activa de la comunidad y aliados.",
        "Establezca sistemas de monitoreo y evaluación continua en entorno real.",
        "Documente casos de éxito y aprendizajes durante la implementación.",
    ],
    7: [
        "Realice ajustes basados en el análisis sistemático de datos y retroalimentación.",
        "Valide las mejoras introducidas con la comunidad objetivo.",
        "Documente un modelo operativo actualizado, claro y viable.",
    ],
    8: [
        "Desarrolle un plan integral para la integración comunitaria e institucional.",
        "Capacite a la comunidad para el mantenimiento y gestión autónoma.",
        "Establezca alianzas estratégicas y respaldo normativo con actores locales.",
    ],
    9: [
        "Consolide mecanismos de sostenibilidad y gobernanza a largo plazo.",
        "Documente historias de éxito y sistematización de lecciones aprendidas.",
        "Establezca protocolos y guías para la replicación en otras comunidades.",
    ],
}

GLOSSARY = [
    {
        "term": "Nivel de preparación social",
        "aliases": ["SRL", "Societal Readiness Level", "madurez social", "preparación social"],
        "definition": "Escala de 9 niveles propuesta por Innovation Fund Denmark que estima el grado de preparación de un resultado o solución para ser adoptado, integrado y sostenido por la sociedad.",
    },
    {
        "term": "Innovación social",
        "aliases": ["innovación social"],
        "definition": "Nuevas ideas, modelos, productos, servicios o procesos que resuelven necesidades sociales de forma más eficaz que las alternativas existentes y generan nuevas relaciones o colaboraciones sociales.",
    },
    {
        "term": "Partes interesadas",
        "aliases": ["partes interesadas", "actores clave", "grupos de interés", "actores locales"],
        "definition": "Personas, comunidades, líderes sociales, organizaciones o instituciones que se ven afectadas por el problema social o que participan en el diseño, validación y adopción de la solución.",
    },
    {
        "term": "Comunidad objetivo",
        "aliases": ["comunidad objetivo", "comunidades objetivo", "personas destinatarias", "persona destinataria", "persona usuaria"],
        "definition": "Grupo poblacional o territorio directamente beneficiario o participante en el que se busca resolver el problema social y consolidar la apropiación de la solución.",
    },
    {
        "term": "Prueba conceptual",
        "aliases": ["prueba conceptual", "versión inicial"],
        "definition": "Primera materialización estructurada de la solución social que permite someter sus supuestos a discusión y validación temprana con las partes interesadas antes de realizar un piloto en terreno.",
    },
    {
        "term": "Prueba piloto",
        "aliases": ["prueba piloto", "prototipar"],
        "definition": "Implementación experimental a escala acotada en un entorno relevante para evaluar el comportamiento real de la solución, recoger datos de impacto potencial e identificar barreras.",
    },
    {
        "term": "Entorno relevante",
        "aliases": ["entorno relevante", "entorno real"],
        "definition": "Contexto comunitario, educativo, organizacional o territorial cuyas condiciones sociales, culturales y operativas reproducen aquellas en las que operará definitivamente la solución.",
    },
    {
        "term": "Modelo operativo",
        "aliases": ["modelo operativo"],
        "definition": "Definición estructurada de los procesos, roles, recursos, canales de relación y mecanismos de gobernanza necesarios para que la solución funcione de manera regular y sostenible.",
    },
    {
        "term": "Replicabilidad",
        "aliases": ["replicabilidad", "replicar", "replique"],
        "definition": "Capacidad de una solución social demostrada para adaptarse y reproducirse con efectividad en otras comunidades, regiones o grupos poblacionales con problemáticas similares.",
    },
    {
        "term": "Cambio sistémico",
        "aliases": ["cambio sistémico", "cambio estructural", "transformación social"],
        "definition": "Transformación duradera en las causas profundas de un problema social, en las relaciones de poder, en la cultura, en el tejido relacional del ecosistema o en las políticas públicas.",
    },
    {
        "term": "Innovación disruptiva",
        "aliases": ["disruptiva", "disruptivo", "incremental", "mejoras incrementales"],
        "definition": "Cambio cualitativo que modifica las bases o el paradigma con que se aborda un problema social, en contraste con la innovación incremental que introduce mejoras parciales sobre el modelo existente.",
    },
    {
        "term": "Co-creación",
        "aliases": ["co-creación", "diagnóstico participativo"],
        "definition": "Proceso de diseño y construcción conjunta en el que investigadores, ciudadanía y actores del territorio participan activamente en la definición del problema y la generación de la solución.",
    },
    {
        "term": "Empoderamiento",
        "aliases": ["empoderamiento", "autonomía"],
        "definition": "Proceso mediante el cual las personas y comunidades fortalecen sus capacidades, confianza y protagonismo para tomar decisiones y gestionar autónomamente las soluciones a sus problemas.",
    },
    {
        "term": "Situación de vulnerabilidad",
        "aliases": ["situación de vulnerabilidad", "personas vulnerables"],
        "definition": "Condición de desventaja o riesgo social, económico, funcional o cultural que limita el ejercicio pleno de derechos y el acceso equitativo a oportunidades y servicios.",
    },
    {
        "term": "Ecosistema",
        "aliases": ["ecosistema", "tejido social", "multiagente"],
        "definition": "Conjunto articulado de actores sociales, comunitarios, académicos, gubernamentales y productivos que interactúan en un territorio alrededor de un reto social.",
    },
]


def clean(value):
    if value is None:
        return ""
    text = str(value)
    if "[Comentario encadenado]" in text:
        parts = re.split(r"Comentario:\s*", text, maxsplit=1)
        text = parts[1] if len(parts) > 1 else text
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"\s+([,.;:)])", r"\1", text)
    for wrong, right in ERRATA.items():
        text = text.replace(wrong, right)
    return text


class SRLHtmlParser(HTMLParser):
    """Extrae las 9 secciones y 45 preguntas del archivo Nivel SRL.html."""

    def __init__(self):
        super().__init__()
        self.sections = []
        self._in_h2 = False
        self._in_q = False
        self._buf = []
        self._current_section = None

    def handle_starttag(self, tag, attrs):
        cls = dict(attrs).get("class", "")
        if tag == "h2" and "section-title" in cls:
            self._in_h2 = True
            self._buf = []
        elif tag == "p" and "question-text" in cls:
            self._in_q = True
            self._buf = []

    def handle_endtag(self, tag):
        if tag == "h2" and self._in_h2:
            self._in_h2 = False
            raw = clean("".join(self._buf))
            # Ej: "Nivel 1: Identificación del Problema"
            m = re.match(r"Nivel\s+(\d+)\s*:\s*(.+)", raw, re.I)
            lvl = int(m.group(1)) if m else len(self.sections) + 1
            name = m.group(2).strip() if m else raw
            self._current_section = {
                "level": lvl,
                "key": f"srl-{lvl}",
                "short": name,
                "title": f"Nivel SRL {lvl}: {name}",
                "stage": SRL_STAGES[lvl],
                "description": SRL_LEVEL_DESCRIPTIONS[lvl],
                "recommendations": SRL_RECOMMENDATIONS[lvl],
                "items": [],
            }
            self.sections.append(self._current_section)
        elif tag == "p" and self._in_q:
            self._in_q = False
            raw = clean("".join(self._buf))
            raw = re.sub(r"^\d+\.\s*", "", raw)
            if self._current_section is not None:
                lvl = self._current_section["level"]
                idx = len(self._current_section["items"]) + 1
                global_idx = (lvl - 1) * 5 + (idx - 1)
                self._current_section["items"].append(
                    {
                        "id": f"{lvl}-{idx}",
                        "srl": lvl,
                        "text": raw,
                        "action": SRL_ACTIONS[global_idx],
                    }
                )

    def handle_data(self, data):
        if self._in_h2 or self._in_q:
            self._buf.append(data)


def read_srl(html_path: Path):
    parser = SRLHtmlParser()
    parser.feed(html_path.read_text(encoding="utf-8"))
    if len(parser.sections) != 9:
        raise SystemExit(f"Se esperaban 9 niveles SRL y se encontraron {len(parser.sections)}")
    total_q = sum(len(s["items"]) for s in parser.sections)
    if total_q != 45:
        raise SystemExit(f"Se esperaban 45 preguntas SRL y se encontraron {total_q}")
    return {
        "thresholdPoints": 7,
        "maxPointsPerLevel": 10,
        "thresholdRatio": 0.70,
        "optionValues": {"no": 0, "partial": 1, "yes": 2},
        "optionLabels": {"no": "No", "partial": "Parcialmente", "yes": "Sí"},
        "maturity": {str(k): v for k, v in SRL_MATURITY.items()},
        "sections": parser.sections,
    }


def read_helix(xlsm_path: Path):
    wb = openpyxl.load_workbook(xlsm_path, data_only=False)
    ws_elem = wb["elementos"]
    ws_pond = wb["ponderacion"]
    ws_pred = wb["predeterminadas"]
    ws_pref = wb["preferencias"]
    ws_param = wb["parametros"]
    ws_frases = wb["frases"]
    ws_p1e1 = wb["P1E1"]

    # Valores de respuestas de valoración: [1, 0.75, 0.5, 0.25, 0]
    option_scores = [float(ws_param.cell(r, 3).value) for r in (14, 16, 18, 20, 22)]
    # Factores correctores de preferencias: [2, 1.5, 1, 0.75, 0.2] (celdas B6:B10)
    pref_factors = [
        {"key": "crucial", "label": "Crucial", "factor": float(ws_param.cell(6, 2).value)},
        {"key": "muy_importante", "label": "Muy importante", "factor": float(ws_param.cell(7, 2).value)},
        {"key": "importante", "label": "Importante", "factor": float(ws_param.cell(8, 2).value)},
        {"key": "algo_importante", "label": "Algo importante", "factor": float(ws_param.cell(9, 2).value)},
        {"key": "irrelevante", "label": "Irrelevante", "factor": float(ws_param.cell(10, 2).value)},
    ]

    # 8 criterios de preferencias (filas 7, 12, 17, 22, 27, 32, 37, 42 en hoja preferencias)
    pref_rows = [7, 12, 17, 22, 27, 32, 37, 42]
    preferences = []
    for idx, r in enumerate(pref_rows):
        raw = clean(ws_pref.cell(r, 2).value)
        parts = [p.strip() for p in raw.split("·", 1)]
        preferences.append(
            {
                "id": f"pref-{idx + 1}",
                "index": idx,
                "short": parts[0],
                "description": parts[1] if len(parts) > 1 else raw,
            }
        )

    # Filas de los 50 elementos en la hoja ponderacion (saltando las filas de suma 24, 35, 46, 57)
    pond_elem_rows = (
        list(range(14, 24))
        + list(range(25, 35))
        + list(range(36, 46))
        + list(range(47, 57))
        + list(range(58, 68))
    )

    # 19 ámbitos predeterminados (columnas I=9 a AA=27 en hoja ponderacion; filas 7..43 paso 2 en predeterminadas)
    presets = []
    for idx, col in enumerate(range(9, 28)):
        pred_row = 7 + idx * 2
        name = clean(ws_pred.cell(pred_row, 4).value)
        vec_weights = [float(ws_pond.cell(r, col).value) for r in range(8, 13)]
        elem_weights = [float(ws_pond.cell(r, col).value) for r in pond_elem_rows]
        presets.append(
            {
                "id": f"preset-{idx}",
                "index": idx,
                "name": name,
                "vectorWeights": vec_weights,
                "elementWeights": elem_weights,
                "excelV7ZeroBug": (col == 22),  # En ponderacion!V7 el autor puso 0 en vez de =+predeterminadas!F33
            }
        )

    # 5 vectores
    vector_ShortNames = ["Innovación", "Impacto", "Cambio sistémico", "Participación", "Consistencia"]
    vector_frases_rows = [8, 10, 12, 14, 16]
    vectors = []
    for v_idx in range(5):
        pond_row = 8 + v_idx
        raw_desc = clean(ws_frases.cell(vector_frases_rows[v_idx], 2).value)
        parts = [p.strip() for p in raw_desc.split("·", 1)]
        pref_mask = [1 if ws_pond.cell(pond_row, c).value == 1 else 0 for c in range(29, 37)]
        vectors.append(
            {
                "id": f"V{v_idx + 1}",
                "index": v_idx,
                "code": clean(ws_elem.cell(6, 4 + v_idx).value),
                "short": vector_ShortNames[v_idx],
                "description": parts[1] if len(parts) > 1 else raw_desc,
                "standardWeight": float(ws_pond.cell(pond_row, 9).value),
                "prefMask": pref_mask,
                "items": [],
            }
        )

    # Localizar encabezados de los 50 elementos en frases y P1E1
    frases_headers = [
        r
        for r in range(40, ws_frases.max_row + 1)
        if isinstance(ws_frases.cell(r, 2).value, str)
        and ws_frases.cell(r, 2).value.startswith("=+elementos!B")
    ]
    p1e1_headers = [
        r
        for r in range(90, ws_p1e1.max_row + 1)
        if isinstance(ws_p1e1.cell(r, 2).value, str)
        and ws_p1e1.cell(r, 2).value.startswith("=+frases!B")
    ]

    if len(frases_headers) != 50 or len(p1e1_headers) != 50:
        raise SystemExit(f"Error en encabezados HELIX: frases={len(frases_headers)}, P1E1={len(p1e1_headers)}")

    for idx in range(50):
        v_idx = idx // 10
        within_idx = (idx % 10) + 1
        elem_row = 14 + idx
        pond_row = pond_elem_rows[idx]
        f_row = frases_headers[idx]
        p_row = p1e1_headers[idx]

        code = clean(ws_elem.cell(elem_row, 1).value)
        full_title = clean(ws_elem.cell(elem_row, 2).value)
        parts = [p.strip() for p in full_title.split("·", 1)]
        label = parts[0] if len(parts) > 1 else code
        question = parts[1] if len(parts) > 1 else full_title

        cm = ws_p1e1.cell(p_row, 5).comment
        explanation = clean(cm.text) if cm else ""

        options = []
        for opt_idx, offset in enumerate((2, 4, 6, 8, 10)):
            phrase = clean(ws_frases.cell(f_row + offset, 4).value)
            options.append(
                {
                    "level": opt_idx + 1,
                    "score": option_scores[opt_idx],
                    "text": phrase,
                }
            )

        pref_mask = [1 if ws_pond.cell(pond_row, c).value == 1 else 0 for c in range(29, 37)]
        std_weight = float(ws_pond.cell(pond_row, 9).value)

        vectors[v_idx]["items"].append(
            {
                "id": f"V{v_idx + 1}-{within_idx}",
                "globalIndex": idx,
                "vectorIndex": v_idx,
                "code": code,
                "label": label,
                "question": question,
                "explanation": explanation,
                "standardWeight": std_weight,
                "prefMask": pref_mask,
                "options": options,
            }
        )

    return {
        "baseScore": float(ws_param.cell(24, 3).value),
        "optionScores": option_scores,
        "naLabel": "No puedo valorar este elemento / No aplica en este proyecto",
        "preferenceFactors": pref_factors,
        "preferences": preferences,
        "presets": presets,
        "vectors": vectors,
    }


def main():
    srl_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SRL_HTML
    helix_path = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_HELIX_XLSM

    srl_data = read_srl(srl_path)
    helix_data = read_helix(helix_path)

    guide = [
        "SocialTech (SRL-UPTC) es la herramienta institucional de la Universidad Pedagógica y Tecnológica de Colombia para diagnosticar la preparación social y valorar la innovación social de los resultados de investigación, extensión y transferencia.",
        "Integra dos módulos complementarios: (1) el diagnóstico del Nivel de Preparación Social (Societal Readiness Level, SRL 1 a 9), basado en la escala de Innovation Fund Denmark, que evalúa mediante 45 preguntas el grado de madurez hacia la adopción e integración social y genera una ruta de maduración hasta SRL 9; y (2) la valoración multidimensional de Innovación Social (HELIX), basada en la metodología del Observatorio de la Realidad Social de Navarra y Humani, estructurada en 5 vectores y 50 atributos ponderables según la naturaleza del proyecto.",
        "Toda la información se procesa localmente y se guarda automáticamente en este navegador: ningún dato sobre su proyecto o diagnóstico se envía a servidores externos.",
    ]

    data = {
        "guide": guide,
        "srl": srl_data,
        "helix": helix_data,
        "glossary": GLOSSARY,
    }

    payload = json.dumps(data, ensure_ascii=False, indent=2)
    OUT.write_text(
        "/* Generado por tools/build_data.py a partir de Nivel SRL.html y HELIX (.xlsm). No editar a mano. */\n"
        f"window.SRL_DATA = {payload};\n",
        encoding="utf-8",
    )
    print(
        f"{OUT} escrito: {sum(len(s['items']) for s in srl_data['sections'])} preguntas SRL, "
        f"{sum(len(v['items']) for v in helix_data['vectors'])} elementos HELIX, "
        f"{len(helix_data['presets'])} ámbitos predeterminados, "
        f"{len(GLOSSARY)} términos de glosario"
    )


if __name__ == "__main__":
    main()
